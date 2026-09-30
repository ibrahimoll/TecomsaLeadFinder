"""Read uploaded CSV/XLSX files into header + rows, with safe limits and clear errors."""

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from app.core.exceptions import BusinessRuleError

SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx"}


@dataclass
class ParsedFile:
    columns: list[str]
    rows: list[dict[str, Any]]  # keyed by column name; values are str | int | float | datetime | None


def detect_file_type(filename: str, head: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".xls"):
        raise BusinessRuleError(
            "Legacy .xls files are not supported. In Excel, use File → Save As → Excel Workbook (.xlsx), "
            "or export from Odoo as XLSX or CSV."
        )
    for ext, kind in SUPPORTED_EXTENSIONS.items():
        if lower.endswith(ext):
            if kind == "xlsx" and not head.startswith(b"PK"):
                raise BusinessRuleError("This file has an .xlsx extension but is not a valid Excel workbook.")
            return kind
    raise BusinessRuleError("Unsupported file type. Upload a .csv or .xlsx file exported from Odoo.")


def _clean_header(value: Any, position: int) -> str:
    text = " ".join(str(value).split()) if value is not None else ""
    return text or f"Column {position + 1}"


def _dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for h in headers:
        if h in seen:
            seen[h] += 1
            result.append(f"{h} ({seen[h]})")
        else:
            seen[h] = 1
            result.append(h)
    return result


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return value


def _build(header_row: list[Any], data_rows, max_rows: int) -> ParsedFile:
    columns = _dedupe_headers([_clean_header(v, i) for i, v in enumerate(header_row)])
    rows: list[dict[str, Any]] = []
    for raw in data_rows:
        values = [_cell(v) for v in raw]
        if not any(v is not None for v in values):
            continue  # skip blank lines
        if len(rows) >= max_rows:
            raise BusinessRuleError(f"The file has more than {max_rows:,} rows. Split it into smaller files.")
        values += [None] * (len(columns) - len(values))
        rows.append(dict(zip(columns, values[: len(columns)])))
    return ParsedFile(columns=columns, rows=rows)


def _parse_csv(content: bytes, max_rows: int) -> ParsedFile:
    for encoding in ("utf-8-sig", "cp1256", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise BusinessRuleError("Could not read the CSV text encoding. Export it again as UTF-8.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = next(reader, None)
    if not header or not any(h.strip() for h in header):
        raise BusinessRuleError("The CSV file is empty or has no header row.")
    return _build(header, reader, max_rows)


def _parse_xlsx(content: bytes, max_rows: int) -> ParsedFile:
    from openpyxl import load_workbook  # imported lazily: only needed for XLSX

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise BusinessRuleError(f"The Excel file could not be opened ({exc.__class__.__name__}). "
                                "Re-export it from Odoo or save it again as .xlsx.") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = None
        for row in rows:  # first non-empty row is the header
            if row and any(v not in (None, "") for v in row):
                header = list(row)
                break
        if header is None:
            raise BusinessRuleError("The first worksheet is empty.")
        return _build(header, rows, max_rows)
    finally:
        workbook.close()


def parse_file(content: bytes, file_type: str, max_rows: int) -> ParsedFile:
    parsed = _parse_csv(content, max_rows) if file_type == "csv" else _parse_xlsx(content, max_rows)
    if not parsed.rows:
        raise BusinessRuleError("The file has a header row but no data rows.")
    return parsed

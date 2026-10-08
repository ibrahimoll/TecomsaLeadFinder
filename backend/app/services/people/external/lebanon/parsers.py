"""Parsers: one public document -> Records (records.py), read exactly as published. No network here.

    abl_almanac   ABL "Almanac of Banks in Lebanon" PDF: per bank, Board of Directors, Senior Management, founder,
                  branch managers; official bank list / commercial registry numbers; website
    icc_register  Insurance Control Commission register page: General Manager, Financial Manager per insurer
                  / broker, licence number, website, the entry's own update date
    ministry_xls  Ministry of Industry licensed-factory list (XLS): factory, owner when the entry names one
                  ("X لصاحبها Y" / "Y صاحب X" / a factory licensed in a person's name), website, e-mail domain
    idal          IDAL exporter directories (PDF): company, contact person (a person only with a published title)
    profiles      association profile page (RDCL...): person, association office, biography (company roles are
                  read from the biography at match time, only where it names the company)
    member_list   association member search page (ALI...): member company names (identity only)
    hospital      Syndicate of Hospitals hospital page: CEO, chair, founding date, beds, phone, e-mail, website, address
    gleif         GLEIF LEI records of Lebanon (JSON): legal name, legal form, registry number, LEI, address, reported
                  direct / ultimate parent
    ppa_awards    Public Procurement Authority award results (listing page): winner(s), tender, date, amount: a
                  company SIGNAL (procurement_award), never a person
Organisation records also carry the company FACTS the source states (Record.facts), stored like website facts.
    csv_file      an official / licensed file supplied by an operator (Commercial Register extract, licensed
                  directory): company, person, title, date, url
"""

import csv
import io
import json
import re

from app.services.normalization import company_domain
from app.services.people.external.lebanon.records import Record

_HONORIFIC = re.compile(r"^(?:(?:H\.?\s?E\.?|Mr\.?|Mrs\.?|Ms\.?|Miss|Dr\.?|Me\.?|Eng\.?|Prof\.?|Sheikh|Sheikha|Emir|"
                        r"Messrs\.?)\s+)+", re.IGNORECASE)
_WEB = re.compile(r"(?:https?://)?(?:www\.)[\w.-]+\.[a-z]{2,}|https?://[\w.-]+\.[a-z]{2,}", re.IGNORECASE)
_EMAIL = re.compile(r"[\w.+-]+@([\w-]+(?:\.[\w-]+)+)")


def _domains(text: str) -> list[str]:
    found = [company_domain(u if u.startswith("http") else f"https://{u}") for u in _WEB.findall(text)]
    found += [company_domain(f"https://{d}") for d in _EMAIL.findall(text)]
    return list(dict.fromkeys(d for d in found if d))


def split_name_title(line: str) -> tuple[str, str] | None:
    """'Mr. Mutasim Khodr Mahmassani Board Member & General Manager' -> (name, title): the shortest name (2-5
    words) followed by a recognised title."""
    from app.services.people.external.lebanon.records import best_title
    from app.services.people.names import is_person_name
    from app.services.discovery.geography import _fold
    from app.services.people.names import PARTICLES
    from app.services.people.titles import is_department, is_title_word

    text = " ".join(line.replace(", Esq.", "").split())
    if not _HONORIFIC.match(text):
        return None
    words = _HONORIFIC.sub("", text).split()
    for k in range(2, min(6, len(words))):
        name, title = " ".join(words[:k]), " ".join(words[k:])
        first = words[k]
        if _fold(first) in PARTICLES or re.fullmatch(r"[A-Z]\.", first):
            continue  # "Al", "Bou", "M.": still the name
        starts = is_title_word(first) or first.lower() in _TITLE_START or is_department(first) or (
            first.isupper() and 2 <= len(first) <= 5)  # "AEGM - Chief Operating Officer"
        if starts and is_person_name(name) and (best_title(title) is not None or title.lower() in _BOARD_SEATS):
            return name, title
    return None


# ---- ABL Almanac of Banks -----------------------------------------------------------------------------------------

_TITLE_START = {"deputy", "assistant", "vice", "head", "chief", "country", "board", "delegated", "executive",
                "senior", "managing", "general", "member", "director", "manager", "independent", "non-executive",
                "secretary", "chairman", "chairwoman", "lebanon", "regional", "acting", "group", "first", "central"}
# bare board seats ("Mrs. Nada Aoun Member" under BOARD OF DIRECTORS)
_BOARD_SEATS = {"member", "director", "independent member", "independent director", "non-executive member",
                "non-executive director", "executive member", "executive director"}
_ABL_HEADER = re.compile(r"Official List of Banks", re.IGNORECASE)
# "SHAREHOLDERS Nationality Shares in Capital", "BOARD OF DIRECTORS", "SENIOR MANAGEMENT IN LEBANON"
_SECTION = re.compile(r"^(?:[A-Z][A-Z&/'-]{2,}\s?)*[A-Z]{4,}[A-Z&/' -]*(?:\s+[A-Z][a-z].*)?$")


def _continuation(line: str) -> bool:
    """'Commercial Banking' under 'Mr. X Assistant General Manager, Head of Group': the rest of a wrapped title, not a
    new person, heading, label, page number or address."""
    from app.services.people.names import is_person_name

    if not line or len(line) > 60 or not line[0].isupper() or line.upper() == line or _HONORIFIC.match(line) \
            or _SECTION.match(line) or ":" in line or re.search(r"\d|=====|@|www\.", line):
        return False
    if re.search(r"\b(s\.?a\.?l|s\.?a\.?r\.?l|bank|banque|rep\.|represented|branch|opened|tel|fax|switchboard)\b",
                 line, re.IGNORECASE):
        return False  # a corporate board seat ("Blominvest Bank sal Member"), an address or contact line
    words = line.split()
    return not any(is_person_name(" ".join(words[:k])) for k in (2, 3) if len(words) >= k)  # a name without "Mr."


def abl_almanac(pages: list[tuple[int, str]]) -> list[Record]:
    lines = [(n, " ".join(line.split())) for n, text in pages for line in text.splitlines()]
    lines = [(n, line) for n, line in lines if line]
    records: list[Record] = []
    seats: list[tuple[str, Record]] = []  # (bank, board / founder record)
    foreign: set[str] = set()  # banks headquartered abroad: their board is the foreign head office's
    bank: Record | None = None
    section = None
    for i, (page, line) in enumerate(lines):
        if _ABL_HEADER.match(line) and i:
            name = lines[i - 1][1]
            if i > 1 and re.match(r"\(|(?:AU|AUX|DE|DU|DES|OF|ET|AND)\b", name):
                # "ARAB FINANCE HOUSE S.A.L." / "(ISLAMIC BANK)"; "SOCIETE GENERALE DE BANQUE" / "AU LIBAN S.A.L."
                name = f"{lines[i - 2][1]} {name}"
            bank = Record(org=[name], excerpt=name, page=lines[i - 1][0])
            records.append(bank)
            if m := re.search(r"Official List of Banks N°\s*:?\s*(\d+)", line):
                bank.identifiers["bdl_bank_list"] = m[1]
            if m := re.search(r"Commercial Registry N°\s*:?\s*([\w ]+?)(?:,|$)", line):
                bank.identifiers["registry:LB"] = m[1].strip()
            section = None
            continue
        if bank is None:
            continue  # the association's own front matter
        if m := re.search(r"Website\s*:\s*(\S+)", line):
            bank.domains += _domains(m[1])
        if m := re.search(r"Number of Employees\s*:\s*([\d,]+)", line):
            bank.facts.setdefault("employees", int(m[1].replace(",", "")))
        if m := re.search(r"Operating branches in Lebanon\s*:\s*(\d+)", line):
            bank.facts.setdefault("branches", int(m[1]))
        if m := re.search(r"Switchboard\s*:\s*(\(\d+\)\s*[\d]{5,8})", line):
            bank.facts.setdefault("phone", m[1])
        if m := re.search(r"E-Mail\s*:\s*([\w.+-]+@[\w.-]+\.\w+)", line):
            bank.facts.setdefault("email", m[1])
        if _SECTION.match(line) and not _HONORIFIC.match(line):
            if re.search(r"\bIN LEBANON\b|HEAD OFFICE IN\b", line):
                foreign.add(bank.org[0])  # "MANAGEMENT IN LEBANON", "HEAD OFFICE IN EGYPT"
            section = ("board" if "BOARD OF DIRECTORS" in line else "management" if "MANAGEMENT" in line
                       else "founder" if line.strip().startswith("FOUNDER") else None)
            continue
        if m := re.search(r"Branch Manager\s*:\s*(.+)$", line):
            name = _HONORIFIC.sub("", re.split(r"\s+(?:E-?Mail|Tel|Fax|Switchboard|Alternate|Opened)\b", m[1])[0])
            name = name.strip(" ,;")
            records.append(Record(org=bank.org, person=name, title="Branch Manager", excerpt=line, page=page))
            continue
        if section == "founder" and _HONORIFIC.match(line):
            founder = _HONORIFIC.sub("", line.replace(", Esq.", "")).strip()
            seats.append((bank.org[0], Record(org=bank.org, person=founder, title="Founder",
                                              status="unknown",  # founding is history, not a current position
                                              excerpt=f"{bank.org[0]} — Founder: {line}", page=page)))
            continue
        if section in ("board", "management") and i + 1 < len(lines) and _HONORIFIC.match(line) and (
                re.search(r"(&|-|–|,|\band)$", line) or _continuation(lines[i + 1][1])):
            joined = f"{line} {lines[i + 1][1]}"  # a title wrapped onto the next line
            line = joined if split_name_title(joined) else line  # too long to read joined: the first line alone
        if section in ("board", "management") and (found := split_name_title(line)):
            name, title = found
            bare = title.lower().strip()
            if section == "board" and bare.split()[-1] in ("member", "director") and len(bare.split()) <= 2:
                # a board seat ("Member", "Independent Director"), not a department director
                title = " ".join([*title.split()[:-1], "Board Member"])
            heading = "Board of Directors" if section == "board" else "Senior Management"
            record = Record(org=bank.org, person=name, title=title, excerpt=f"{bank.org[0]} — {heading}: {line}",
                            page=page)
            if section == "board":
                seats.append((bank.org[0], record))
            else:
                records.append(record)
    # A foreign bank's branch lists its head office's board / founders: not people of the Lebanese entity.
    return records + [r for name, r in seats if name not in foreign]


def abl_hr(entries: list[dict]) -> list[Record]:
    """ABL 'Heads of HR': [{MemberName: 'FRANSABANK S.A.L. (1)', PersonName: 'Ms. Rania Jamal', MemberSite}]."""
    records = []
    for e in entries:
        bank = re.sub(r"\s*\(\d+\)\s*$", "", (e.get("MemberName") or "").strip())
        person = _HONORIFIC.sub("", (e.get("PersonName") or "").strip())
        if bank and person:
            records.append(Record(org=[bank], person=person, title="Head of HR",
                                  domains=_domains(e.get("MemberSite") or ""),
                                  excerpt=f"Association of Banks in Lebanon - Heads of HR: {e.get('MemberName')} - "
                                          f"{e.get('PersonName')}"))
    return records


# ---- ICC register ------------------------------------------------------------------------------------------------

def icc_register(html: str) -> tuple[list[Record], int]:
    """(records, number of pages of the register)."""
    m = re.search(r"window\.__remixContext\s*=\s*(\{.*?\});\s*</script>", html, re.S)
    if not m:
        return [], 0
    data = json.loads(m[1])["state"]["loaderData"]
    route = next((v for k, v in data.items() if k.endswith("companies") and isinstance(v, dict)), {})
    records = []
    for entry in route.get("companies") or []:
        a = entry.get("attributes") or {}
        name, kind = (a.get("Name") or "").strip(), a.get("Type") or ""
        if not name:
            continue
        org = [p.strip() for p in name.split("|") if p.strip()]
        base = dict(org=org, domains=_domains(a.get("Website") or ""),
                    identifiers={"icc_licence": str(a["Number"])} if a.get("Number") else {},
                    published=(a.get("updatedAt") or "")[:10] or None)
        phone = re.split(r"\s+[-/]?\s*(?=\+)", (a.get("Phone") or "").strip())[0]
        facts = {k: v for k, v in {"phone": phone, "address": (a.get("Location") or "").strip(),
                                   "location_type": "headquarters"}.items() if v}
        records.append(Record(**base, excerpt=f"{name} — {kind}, ICC register", facts=facts))
        for field, title in (("GeneralManger", "General Manager"), ("GeneralManager", "General Manager"),
                             ("FinancialManager", "Financial Manager")):
            if person := (a.get(field) or "").strip():
                records.append(Record(**base, person=person, title=title,
                                      excerpt=f"{name} — {kind}. {title}: {person} (Insurance Control Commission "
                                              f"register, licence {a.get('Number') or 'n/a'})"))
    return records, int(route.get("maxPages") or 1)


# ---- Ministry of Industry licensed factories ---------------------------------------------------------------------

_OWNED_BY = re.compile(r"^(?P<org>.+?)\s*\(?\s*لصاحب(?:ها|ه|تها|يها)\s+(?P<person>[^)()\-–]+?)\s*\)?\s*(?:-.*)?$")
_OWNER_OF = re.compile(r"^(?:السيد|السيدة)?\s*(?P<person>[^()]+?)\s+صاحب(?:ة)?\s+(?:المحل التجاري|محلات|شركة|"
                       r"مؤسسة|المؤسسة|معمل|مصنع)?\s*[\"«]?(?P<org>.+?)[\"»]?\s*\)?$")
# "Y المعروف بالاسم التجاري Z" / "Y بالإسم التجاري Z": person Y trading as Z
_TRADE_NAME = re.compile(r"^(?:السيد|السيدة)?\s*(?P<person>[^()]+?)\s+(?:المعروف\s+)?بال[اإ]سم التجاري\s+(?P<org>.+)$")
_PERSON_LICENCE = re.compile(r"^(?:السيد|السيدة)\s+(?P<person>[^()/\-–]+)$")


def ministry_rows(rows: list[list[str]], year: int | None) -> list[Record]:
    """rows: [ID, factory / owner, governorate, caza, town, district, street, building, website, phones...,
    email, production] as in the Ministry's lists."""
    records = []
    for row in rows:
        if len(row) < 17 or not str(row[1]).strip():
            continue
        raw = " ".join(str(row[1]).replace('"', " ").replace("''", " ").split())
        ident = str(row[0]).split(".")[0]
        place = " - ".join(str(x).strip() for x in (row[3], row[4]) if str(x).strip())
        common = dict(domains=_domains(f"{row[8]} {row[15]}"),
                      identifiers={"moi_factory": ident} if ident.isdigit() else {},
                      published=f"{year}-12-31" if year else None, page=None)
        excerpt = f"{raw} — {place}: {str(row[16]).strip()} (Ministry of Industry, licensed factories {year or ''})"
        person = org = None
        if m := _TRADE_NAME.match(raw):
            org, person = m["org"].strip(" -\"'"), m["person"].strip()
        elif m := _OWNED_BY.match(raw):
            org, person = m["org"].strip(" -"), m["person"].strip()
        elif m := _OWNER_OF.match(raw):
            org, person = m["org"].strip(" -"), m["person"].strip()
        elif m := _PERSON_LICENCE.match(raw):
            person = m["person"].strip()
            org = person  # a factory licensed in its owner's own name
        if person and " و" in f" {person}":
            person = None  # "حسن وطلال": several owners in one phrase, not read
        phone = next((str(p).strip() for p in row[9:14] if str(p).strip()), None)
        email = str(row[15]).strip()
        facts = {"address": " - ".join(str(x).strip() for x in (row[6], row[5], row[4], row[3]) if str(x).strip()),
                 "city": str(row[4]).strip() or None, "location_type": "factory",
                 "signals": ["manufacturing", "factory"]}
        if phone:
            facts["phone"] = phone
        if "@" in email:
            facts["email"] = email
        records.append(Record(org=[org or raw], excerpt=excerpt[:500], facts=facts, **common))
        if person:
            records.append(Record(org=[org or raw], person=person, title="Owner", excerpt=excerpt[:500], **common))
    return records


def ministry_xls(raw: bytes, year: int | None) -> list[Record]:
    if raw[:4] == b"PK\x03\x04":  # .xlsx
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(raw), read_only=True).worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in sheet.iter_rows(min_row=2, values_only=True)]
    else:
        import xlrd

        sheet = xlrd.open_workbook(file_contents=raw).sheet_by_index(0)
        rows = [[str(c.value) for c in sheet.row(i)] for i in range(1, sheet.nrows)]
    return ministry_rows(rows, year)


# ---- IDAL directories ------------------------------------------------------------------------------------------

def idal(pages: list[tuple[int, str]], year: int | None) -> list[Record]:
    lines = [(n, " ".join(line.split())) for n, text in pages for line in text.splitlines()]
    lines = [(n, line) for n, line in lines if line and not re.fullmatch(r"[\d\s]+|IDAL \| \d{4}", line)]
    starts = [i for i in range(len(lines) - 1) if not lines[i][1].startswith("•") and (
        lines[i + 1][1].startswith("•") or lines[i + 1][1] == "Company Profile")]
    records = []
    for k, i in enumerate(starts):
        page, name = lines[i]
        block = [line for _, line in lines[i:(starts[k + 1] if k + 1 < len(starts) else len(lines))]]
        text = "\n".join(block)
        base = dict(org=[name], domains=_domains(text), page=page,
                    published=f"{year}-12-31" if year else None)
        facts: dict = {"signals": ["exporting"]}
        if m := re.search(r"(?:Number of Employees|No\. of Employees)\s*:?\s*\n?\s*([\d,]+)", text):
            facts["employees"] = int(m[1].replace(",", ""))
        if m := re.search(r"Phone\s*:?\s*\n?\s*([\d][\d ]{7,16}\d)", text):
            facts["phone"] = "+" + re.sub(r"\D", "", m[1]) if m[1].startswith("961") else m[1]
        if m := re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text):
            facts["email"] = m[0]
        records.append(Record(**base, excerpt=name, facts=facts))
        contact = None
        for j, line in enumerate(block):
            if m := re.match(r"•?\s*Contact(?: Person)?\s*:\s*(.+)$", line):
                contact = m[1]
            elif line == "Contact Person" and j + 1 < len(block):
                contact = block[j + 1]
        if contact and "," in contact:  # "Mr. Nazih Tabet, CEO": a person only with a published title
            person, title = (x.strip() for x in contact.split(",", 1))
            records.append(Record(**base, person=_HONORIFIC.sub("", person), title=title,
                                  excerpt=f"{name} — Contact Person: {contact} (IDAL directory {year or ''})"))
    return records


# ---- association pages -------------------------------------------------------------------------------------------

def profile(doc, association: str) -> Record | None:
    """An association's profile page of one person: name (page heading), office at the association, biography."""
    heading = next((t for tag, t, area in doc.blocks if tag in ("h1", "h2") and area != "nav"), None) or \
        (doc.title or "").split(" - ")[0]
    texts = [t for tag, t, area in doc.blocks if area not in ("nav", "footer") and t != heading]
    bio = " ".join(t for t in texts if len(t) > 80)
    if not heading or not bio:
        return None
    office = next((t for t in texts if 2 <= len(t) <= 60), None)
    return Record(org=[], person=heading.strip(), association_role=office, bio=bio[:4000],
                  excerpt=f"{association} profile: {heading}")


def member_list(html: str) -> list[Record]:
    """ALI 'Find a member' results: member company names (and ALI's own member reference), identity only."""
    records = []
    for href, name in re.findall(r'<a[^>]+href="([^"]*)"[^>]*class="member-title"[^>]*>\s*<p>(.*?)</p>', html, re.S):
        name = " ".join(re.sub(r"<[^>]+>", " ", name).split())
        member = re.search(r"post=(\d+)", href)
        if name:
            records.append(Record(org=[name], excerpt=f"Association of Lebanese Industrialists member: {name}",
                                  identifiers={"ali_member": member[1]} if member else {}))
    return records


# ---- Syndicate of Hospitals ---------------------------------------------------------------------------------------

_HOSPITAL_FIELDS = ("Date of Establishment", "Establisher", "Chair Person", "CEO", "Location", "Address",
                    "Phone Number(s)", "Fax", "E-mail", "Web Site", "Accredited by the Lebanese Ministry of Public Health",
                    "Type of Care Offered", "Number of Beds")


def hospital(doc) -> list[Record]:
    """A hospital's page in the Syndicate of Hospitals directory: 'Label | value' pairs."""
    texts = [t.strip() for _, t, area in doc.blocks if area not in ("nav", "footer") and t.strip()]
    values: dict[str, str] = {}
    for i, t in enumerate(texts[:-1]):
        if t in _HOSPITAL_FIELDS and t not in values and texts[i + 1] not in _HOSPITAL_FIELDS:
            values[t] = texts[i + 1]
    name = next((t for tag, t, area in doc.blocks if tag in ("th", "h1", "h2", "h3") and area != "nav"
                 and t.strip() not in _HOSPITAL_FIELDS), None)
    if not name or not values:
        return []
    facts: dict = {"signals": ["healthcare", "hospitals"], "location_type": "headquarters"}
    if address := values.get("Address"):
        facts["address"] = address
    if city := values.get("Location"):
        facts["city"] = city
    if phone := values.get("Phone Number(s)"):
        facts["phone"] = re.split(r"\s*[/,]\s*", phone)[0]
    if "@" in (email := values.get("E-mail", "")):
        facts["email"] = email.split()[0]
    if m := re.search(r"(1[89]\d\d|20[0-2]\d)", values.get("Date of Establishment", "")):
        facts["founded"] = int(m[1])
    if (beds := values.get("Number of Beds", "")).isdigit():
        facts["beds"] = int(beds)
    base = dict(org=[name.strip()], domains=_domains(values.get("Web Site", "")))
    records = [Record(**base, excerpt=f"Syndicate of Hospitals in Lebanon: {name.strip()}", facts=facts)]
    for label, title in (("CEO", "CEO"), ("Chair Person", "Chairperson")):
        if values.get(label):  # an organisation as chair ("Saint Joseph University") is not a person: dropped later
            person = _HONORIFIC.sub("", re.sub(r"^(?:The\s+)?(?:Reverend|Rev\.|Father|Sister|Pr\.?)\s+", "",
                                              values[label])).strip()
            records.append(Record(**base, person=person, title=title,
                                  excerpt=f"{name.strip()} — {label}: {values[label]} (Syndicate of Hospitals in "
                                          f"Lebanon, hospital directory)"))
    return records


# ---- GLEIF ---------------------------------------------------------------------------------------------------------

def gleif(entries: list[dict]) -> list[Record]:
    """GLEIF LEI records (API 'data' items, each with optional 'parents': [{type, lei, name}])."""
    records = []
    for item in entries:
        a = item.get("attributes") or {}
        e = a.get("entity") or {}
        if e.get("status") != "ACTIVE":
            continue
        names = [(e.get("legalName") or {}).get("name")] + [n.get("name") for n in e.get("otherNames") or []] + \
            [n.get("name") for n in e.get("transliteratedOtherNames") or []]
        names = [n for n in names if n]
        if not names:
            continue
        address = e.get("legalAddress") or {}
        facts = {"legal_name": names[0], "legal_form": (e.get("legalForm") or {}).get("other"),
                 "address": ", ".join(address.get("addressLines") or []), "city": address.get("city"),
                 "location_type": "headquarters",
                 "relationships": [{"type": p["type"], "name": p["name"], "lei": p["lei"]}
                                   for p in item.get("parents") or []]}
        identifiers = {"lei": a.get("lei")}
        if e.get("registeredAs"):
            identifiers["registry:LB"] = str(e["registeredAs"])
        records.append(Record(org=names, identifiers={k: v for k, v in identifiers.items() if v},
                              excerpt=f"GLEIF LEI {a.get('lei')}: {names[0]} ({facts['legal_form'] or 'legal entity'})"
                                      f", {address.get('city') or ''}", facts={k: v for k, v in facts.items() if v},
                              published=(a.get("registration") or {}).get("lastUpdateDate", "")[:10] or None))
    return records


# ---- Public Procurement Authority awards ----------------------------------------------------------------------------

def _cell(html: str) -> str:
    return " ".join(re.sub(r"<[^>]+>", " ", html).split())


def ppa_awards(html: str) -> list[Record]:
    """Rows of www.ppa.gov.lb/ar/awards: procuring entity (logo), tender (link), result date, standstill end,
    winner(s) (one <p> each), amount, details."""
    records = []
    body = html[html.find("<tbody>"):]
    for row in re.findall(r"<tr>(.*?)</tr>", body, re.S):
        cells = re.findall(r"<td[^>]*>(.*?)</td>", row, re.S)
        if len(cells) < 6:
            continue
        link = re.search(r'href="([^"]*/tenders/details/(\d+)/awards)"', cells[1])
        winners = [_cell(p) for p in re.findall(r"<p[^>]*>(.*?)</p>", cells[4], re.S)] or [_cell(cells[4])]
        amount = _cell(cells[5])
        date = _cell(cells[2])[:10]
        title = _cell(cells[1]).split(" / ")[0].strip(" /-")
        if not link:
            continue
        for winner in filter(None, winners):
            award = {"tender_id": link[2], "url": link[1], "title": title, "date": date, "amount": amount}
            records.append(Record(org=[winner], published=date or None, facts={"award": award},
                                  excerpt=f"Public Procurement Authority award ({date}): {winner} — {title or 'tender'}"
                                          f" — {amount}"))
    return records


# ---- Lebanese Contractors Syndicate guide -------------------------------------------------------------------------

def lcs_contractor(html: str, published: str | None = "2020-12-31") -> list[Record]:
    """One member page of the contractors' guide (fields by their stable element ids). The guide is not dated: its
    people are kept 'not confirmed' (as of its last known edition), never current."""
    fields = {k: " ".join(v.replace("ـ", "").split()) for k, v in re.findall(
        r'id="FormView1_(\w+)"[^>]*>([^<]*)</span>', html or "") if v.strip() and not k.startswith("Label")}
    company = fields.get("MemberName0", "").strip(" -")
    if not company:
        return []
    email = next((e for e in re.split(r"\s*-\s*|\s+", fields.get("Email0", "")) if "@" in e), None)
    first = re.match(r"\s*(\d{2}\s*/\s*\d{6,7})", fields.get("OfficeTel5", ""))  # '01/857357-853749'
    phone = first[1].replace(" ", "") if first else ""
    address = " - ".join(x for x in (fields.get("Street0"), fields.get("Area0"), fields.get("City0")) if x)
    facts = {"signals": ["construction"], "location_type": "office", "city": fields.get("City0")}
    if phone:
        facts["phone"] = phone
    if email:
        facts["email"] = email
    if address:
        facts["address"] = address
    base = dict(org=[company], identifiers={"lcs_member": fields["JoinNumber0"]} if fields.get("JoinNumber0") else {},
                published=published)
    excerpt = f"{company} — Lebanese Contractors Syndicate guide, member {fields.get('JoinNumber0', '?')}"
    records = [Record(**base, excerpt=excerpt, facts=facts)]
    for key, label in (("ManagerName0", "المدير"), ("EngineerName0", "المهندس")):
        if fields.get(key):
            records.append(Record(**base, person=fields[key], title=label, status="unknown",
                                  excerpt=f"{excerpt}: {label} {fields[key]}"))
    return records


# ---- operator-supplied official / licensed files --------------------------------------------------------------------

def csv_file(text: str) -> list[Record]:
    """Columns: company, person, title, published (YYYY-MM-DD), url, domain, registration_number, email, phone."""
    records = []
    for row in csv.DictReader(io.StringIO(text)):
        row = {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
        if not row.get("company"):
            continue
        records.append(Record(
            org=[row["company"]], person=row.get("person") or None, title=row.get("title") or None,
            excerpt=" — ".join(x for x in (row["company"], row.get("person"), row.get("title")) if x),
            domains=[d for d in [company_domain(f"https://{row['domain']}")] if d] if row.get("domain") else [],
            identifiers={"registry:LB": row["registration_number"]} if row.get("registration_number") else {},
            email=row.get("email") or None, phone=row.get("phone") or None,
            published=row.get("published") or None))
    return records

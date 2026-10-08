"""Step 6.1 acquisition: source-to-many tables (exhibitor / member lists, XLSX, CSV), business cards (.vcf), the
company site's contact artifacts and schema.org people with their contact points. Contacts go to a person only
when the entry names the person and the address is not a role mailbox; home numbers are never read."""

import io
import json

from openpyxl import Workbook
from sqlalchemy import select

from app.models import Email, Phone
from app.models.enums import UserRole
from app.services.people.external import generic
from app.services.people.external.lebanon import configured
from tests.test_enrichment import new_company, research, site  # noqa: F401 - the mocked-web fixture

EXHIBITORS = """<table>
<tr><th>Exhibitor</th><th>Contact Person</th><th>Position</th><th>E-mail</th><th>Mobile</th><th>Website</th></tr>
<tr><td>Omega Foods SAL</td><td>Rana Saad</td><td>Procurement Manager</td><td>rana.saad@omegafoods.example</td>
    <td>03 123 456</td><td>www.omegafoods.example</td></tr>
<tr><td>Delta Trading</td><td>Karim Aoun</td><td>IT Director</td><td>info@delta.example</td><td>71 222 333</td><td></td></tr>
<tr><td>Sigma Hotels</td><td>Lina Khoury</td><td></td><td>lina@sigma.example</td><td></td><td></td></tr>
</table>"""


def test_exhibitor_table_source_to_many_with_contact_ownership():
    records = configured.tables(EXHIBITORS, {"event": {"name": "HORECA Lebanon", "year": 2026}})
    people = {r.person: r for r in records if r.person}
    assert set(people) == {"Rana Saad", "Karim Aoun"}  # Lina has no title: never a person record
    rana, karim = people["Rana Saad"], people["Karim Aoun"]
    assert (rana.email, rana.phone, rana.domains) == ("rana.saad@omegafoods.example", "03 123 456",
                                                     ["omegafoods.example"])
    assert karim.email is None  # info@ is the company's mailbox, never Karim's
    delta = next(r for r in records if r.org == ["Delta Trading"] and not r.person)
    assert delta.facts["email"] == "info@delta.example" and delta.facts["event"]["name"] == "HORECA Lebanon"
    assert rana.published == "2026-12-31"  # an event's year: dated evidence, never "current" forever


def test_spreadsheet_csv_and_arabic_headers():
    book = Workbook()
    sheet = book.active
    sheet.append(["Members directory 2026"])
    sheet.append(["Company", "Name", "Title", "Email", "Tel"])
    sheet.append(["Omega Foods", "Joseph Haddad", "Chief Executive Officer", "jhaddad@omegafoods.example", "01 800 800"])
    raw = io.BytesIO()
    book.save(raw)
    records = configured.files("https://assoc.example/members.xlsx", raw.getvalue(), {})
    person = next(r for r in records if r.person)
    assert (person.person, person.title, person.email, person.phone) == (
        "Joseph Haddad", "Chief Executive Officer", "jhaddad@omegafoods.example", "01 800 800")
    text = "الشركة,الاسم,المنصب,البريد\nأوميغا,رنا سعد,مديرة المشتريات,rana@omega.example\n"
    rows = configured.files("https://assoc.example/list.csv", text.encode(), {})
    assert any(r.person == "رنا سعد" and r.email == "rana@omega.example" for r in rows)


def test_vcard_business_card_never_home_numbers():
    card = ("BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Rana Saad\r\nN:Saad;Rana;;;\r\nORG:Omega Foods\r\n"
            "TITLE:Procurement Manager\r\nEMAIL;TYPE=WORK:rana.saad@omegafoods.example\r\n"
            "EMAIL;TYPE=HOME:rana@gmail.com\r\nTEL;TYPE=CELL:+961 3 123 456\r\nTEL;TYPE=HOME:+961 1 999 999\r\n"
            "END:VCARD\r\n")
    person = next(r for r in generic.vcard_records(card) if r.person)
    assert (person.org, person.email, person.phone) == (["Omega Foods"], "rana.saad@omegafoods.example",
                                                       "+961 3 123 456")
    assert "999 999" not in json.dumps(person.to_json())  # a home number is never read


HOME = """<html><head><title>Omega Foods</title></head><body>
<a href="/team">Our team</a><a href="/cards/rana-saad.vcf">Rana's card</a>
<a href="/docs/management-team.xlsx">Management team</a><a href="/docs/dealers.xlsx">Dealers</a>
<p>Omega Foods makes snacks in Zahle.</p></body></html>"""
TEAM = """<html><title>Our team</title><body><h1>Management</h1><h3>Joseph Haddad</h3><p>Chief Executive Officer</p>
<script type="application/ld+json">{"@context": "https://schema.org", "@type": "Person", "name": "Karim Aoun",
"jobTitle": "IT Manager", "worksFor": {"@type": "Organization", "name": "Omega Foods"},
"email": "mailto:karim.aoun@omegafoods.example", "telephone": "+961 3 444 555"}</script></body></html>"""
CARD = ("BEGIN:VCARD\nVERSION:3.0\nFN:Rana Saad\nORG:Omega Foods\nTITLE:Procurement Manager\n"
        "EMAIL:rana.saad@omegafoods.example\nTEL;TYPE=CELL:+961 3 123 456\nEND:VCARD\n")


def _xlsx(rows) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


def test_site_cards_documents_and_structured_people(client_for, db, site):  # noqa: F811
    agent, _ = client_for(UserRole.SALES_AGENT)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    site.sites["www.omegafoods.example"] = {
        "/": HOME, "/team": TEAM,
        "/cards/rana-saad.vcf": (200, {"content-type": "text/vcard"}, CARD.encode()),
        "/docs/management-team.xlsx": (200, {"content-type": xlsx}, _xlsx([
            ["Name", "Position", "Email", "Mobile"],
            ["Lina Khoury", "Finance Director", "lina.khoury@omegafoods.example", "70 111 222"]])),
        "/docs/dealers.xlsx": (200, {"content-type": xlsx}, _xlsx([
            ["Company", "Name", "Position", "Mobile"], ["Beta Retail", "Tony Rizk", "Owner", "03 999 000"]]))}
    company_id = new_company(agent, name="Omega Foods", website="https://www.omegafoods.example")
    research(agent, company_id)
    people = {p["person"]["full_name"]: p for p in agent.get(f"/companies/{company_id}/people").json()}
    assert {"Joseph Haddad", "Karim Aoun", "Rana Saad", "Lina Khoury"} <= set(people)
    assert "Tony Rizk" not in people  # a dealer list names another company's people
    assert "/docs/dealers.xlsx" not in " ".join(site.requests)  # not a people document: never fetched
    def contact(name, model, column):
        return [getattr(x, column) for x in db.scalars(select(model).where(
            model.person_id == people[name]["person"]["id"]))]
    assert contact("Rana Saad", Email, "email") == ["rana.saad@omegafoods.example"]
    assert contact("Rana Saad", Phone, "phone_type") == ["business_mobile"]
    assert contact("Karim Aoun", Email, "email") == ["karim.aoun@omegafoods.example"]  # schema.org Person email
    assert contact("Lina Khoury", Phone, "normalized_phone") == ["+96170111222"]
    assert people["Rana Saad"]["contactable"]
